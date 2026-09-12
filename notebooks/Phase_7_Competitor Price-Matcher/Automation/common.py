import re
import time
import random
from dataclasses import dataclass
from typing import Optional, List, Literal
from urllib.parse import urlparse
from urllib.robotparser import RobotFileParser

import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIG
# ============================================================

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 "
        "(KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "en-GB,en;q=0.9",
}

REQUEST_TIMEOUT = 15

DELAY_BETWEEN_REQUESTS = (
    1.0,
    2.0,
)

INPUT_FILE = "SOA.xlsx"
OUTPUT_FILE = "price_comparison.xlsx"


# ============================================================
# RETAILERS
# ============================================================


RETAILER_COLUMNS = [
    "Harvey Norman",
    "Currys",
    "Argos",
    "Donaghy Bros",
    "DSE",
]

ALL_COLUMNS = (
    ["Model"]
    + RETAILER_COLUMNS
)


# ============================================================
# FILTER KEYWORDS
# ============================================================

ACCESSORY_KEYWORDS = [
    "cover",
    "case",
    "bag",
    "filter",
    "brush",
    "attachment",
    "spare",
    "replacement part",
    "adapter",
    "cable",
    "charger only",
    "stand for",
    "mat for",
    "accessory",
    "accessories",
]

SPONSORED_KEYWORDS = [
    "sponsored",
    "you may also like",
    "similar item",
    "related product",
]


# ============================================================
# PRICE PARSING
# ============================================================

PRICE_RE = re.compile(
    r"(?:£|GBP\s*)\s*"
    r"(\d+(?:,\d{3})*(?:\.\d{1,2})?)"
    r"(?!\s*/\s?mo|\s*per\s?month|\s*pm)",
    re.IGNORECASE,
)

FINANCE_RE = re.compile(
    r"/\s?mo|per\s?month|\bpm\b|deposit|APR|finance",
    re.IGNORECASE,
)

OUT_OF_STOCK_PHRASES = [
    "out of stock",
    "sold out",
    "unavailable",
    "currently unavailable",
    "not available",
    "no longer available",
    "notify me",
]


# ============================================================
# PRODUCT RESULT
# ============================================================

@dataclass
class ProductResult:
    title: str
    price: Optional[float]
    availability: str
    url: str


# ============================================================
# GENERAL HELPERS
# ============================================================

def polite_sleep():
    """
    Random delay between requests.
    """
    time.sleep(
        random.uniform(
            *DELAY_BETWEEN_REQUESTS
        )
    )


def normalize(text: str) -> str:
    """
    Normalise model strings for comparison.

    Example:
    ES-601 UK -> ES601UK
    """
    return re.sub(
        r"[\s\-_]",
        "",
        text or ""
    ).upper()


def title_contains_model(
    title: str,
    model: str
) -> bool:
    """
    Check whether a product title contains the model number.
    """
    if not title or not model:
        return False

    return (
        normalize(model)
        in normalize(title)
    )


def looks_like_accessory(
    title: str
) -> bool:
    """
    Reject obvious accessories.
    """
    text = (
        title or ""
    ).lower()

    return any(
        keyword in text
        for keyword in ACCESSORY_KEYWORDS
    )


def looks_sponsored(
    text: str
) -> bool:
    """
    Reject sponsored / related-product blocks.
    """
    text = (
        text or ""
    ).lower()

    return any(
        keyword in text
        for keyword in SPONSORED_KEYWORDS
    )


# ============================================================
# PRICE / AVAILABILITY EXTRACTION
# ============================================================

def extract_price(
    text: str
) -> Optional[float]:
    """
    Extract a GBP price from text.

    Ignores obvious finance/monthly-payment lines.
    """

    if not text:
        return None

    cleaned_lines = []

    for line in re.split(
        r"[\n\r]+",
        text
    ):
        if not FINANCE_RE.search(line):
            cleaned_lines.append(line)

    cleaned = " ".join(
        cleaned_lines
    )

    match = PRICE_RE.search(
        cleaned
    )

    if not match:
        return None

    try:
        return float(
            match.group(1).replace(
                ",",
                ""
            )
        )
    except ValueError:
        return None


def extract_availability(
    text: str
) -> str:
    """
    Return IN_STOCK or OUT_OF_STOCK.
    """

    text = (
        text or ""
    ).lower()

    if any(
        phrase in text
        for phrase in OUT_OF_STOCK_PHRASES
    ):
        return "OUT_OF_STOCK"

    return "IN_STOCK"


def best_match(
    candidates: List[ProductResult]
) -> Optional[ProductResult]:
    """
    Prefer the cheapest in-stock product that has a price.
    """

    if not candidates:
        return None

    valid = [
        candidate
        for candidate in candidates
        if (
            candidate.availability
            == "IN_STOCK"
            and candidate.price
            is not None
        )
    ]

    if valid:
        return min(
            valid,
            key=lambda item: item.price
        )

    return candidates[0]


# ============================================================
# OPTIONAL ROBOTS / REQUEST HELPERS
#
# These remain available for retailers where requests +
# BeautifulSoup are useful.
#
# Harvey Norman will later use Playwright directly instead.
# ============================================================

RobotsStatus = Literal[
    "ALLOWED",
    "DISALLOWED",
    "UNAVAILABLE",
]


def robots_status(
    url: str,
    session: requests.Session,
    user_agent: str = "*",
) -> RobotsStatus:
    """
    Check robots.txt.

    Returns:
        ALLOWED
        DISALLOWED
        UNAVAILABLE
    """

    parsed = urlparse(url)

    robots_url = (
        f"{parsed.scheme}://"
        f"{parsed.netloc}/robots.txt"
    )

    parser = RobotFileParser()

    parser.set_url(
        robots_url
    )

    try:
        response = session.get(
            robots_url,
            headers=HEADERS,
            timeout=REQUEST_TIMEOUT,
        )

        response.raise_for_status()

        parser.parse(
            response.text.splitlines()
        )

    except requests.RequestException as exc:
        print(
            f"    [warn] could not check "
            f"{robots_url}: {exc}"
        )

        return "UNAVAILABLE"

    if parser.can_fetch(
        user_agent,
        url
    ):
        return "ALLOWED"

    return "DISALLOWED"


def fetch(
    url: str,
    session: requests.Session,
):
    """
    Fetch a page using requests.

    Returns:
        (BeautifulSoup, "OK")
        (None, "BLOCKED_BOT_DETECTION")
        (None, "ERROR")
    """

    try:
        response = session.get(
            url,
            headers=HEADERS,
            timeout=REQUEST_TIMEOUT,
        )

    except requests.RequestException as exc:
        print(
            f"    [warn] request failed: "
            f"{exc}"
        )

        return None, "ERROR"

    if response.status_code == 200:
        return (
            BeautifulSoup(
                response.text,
                "html.parser"
            ),
            "OK",
        )

    if response.status_code in (
        403,
        429,
        503,
    ):
        print(
            f"    [warn] possible bot block - "
            f"HTTP {response.status_code}"
        )

        return (
            None,
            "BLOCKED_BOT_DETECTION",
        )

    print(
        f"    [warn] unexpected HTTP "
        f"status {response.status_code}"
    )

    return None, "ERROR"


# ============================================================
# EXCEL INPUT
# ============================================================

def read_models(
    input_path: str = INPUT_FILE
) -> List[str]:
    """
    Read model numbers from column A of SOA.xlsx.

    Row 1 is assumed to be a heading.
    """

    workbook = openpyxl.load_workbook(
        input_path,
        data_only=True,
    )

    worksheet = workbook.active

    models = []

    for row in worksheet.iter_rows(
        min_row=2,
        values_only=True,
    ):
        if not row:
            continue

        value = row[0]

        if value is None:
            continue

        model = str(
            value
        ).strip()

        if model:
            models.append(model)

    workbook.close()

    return models


# ============================================================
# EXCEL STYLING
# ============================================================

def _style_header(ws):
    """
    Style the output workbook header.
    """

    header_fill = PatternFill(
        start_color="1F4E78",
        end_color="1F4E78",
        fill_type="solid",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for index in range(
        1,
        len(ALL_COLUMNS) + 1
    ):
        cell = ws.cell(
            row=1,
            column=index,
        )

        cell.fill = header_fill
        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center"
        )


# ============================================================
# EXCEL FORMULAS
# ============================================================

def _add_formulas(
    ws,
    row_num: int
):
    """
    Add Best Price and Best Retailer formulas.

    Automatically uses all retailers listed in RETAILER_COLUMNS.
    """

    if (
        "Best Price" not in ALL_COLUMNS
        or "Best Retailer" not in ALL_COLUMNS
    ):
        return

    col_index = {
        name: index + 1
        for index, name
        in enumerate(ALL_COLUMNS)
    }

    first_retailer = (
        RETAILER_COLUMNS[0]
    )

    last_retailer = (
        RETAILER_COLUMNS[-1]
    )

    first_col = get_column_letter(
        col_index[first_retailer]
    )

    last_col = get_column_letter(
        col_index[last_retailer]
    )

    retailer_range = (
        f"{first_col}{row_num}:"
        f"{last_col}{row_num}"
    )

    retailer_header_range = (
        f"{first_col}$1:"
        f"{last_col}$1"
    )

    best_price_col = (
        get_column_letter(
            col_index["Best Price"]
        )
    )

    best_price_cell = ws.cell(
        row=row_num,
        column=col_index["Best Price"],
    )

    best_price_cell.value = (
        f'=IFERROR('
        f'IF(COUNT({retailer_range})=0,'
        f'"N/A",'
        f'MIN({retailer_range})),'
        f'"N/A")'
    )

    best_price_cell.number_format = (
        '£#,##0.00'
    )

    best_retailer_cell = ws.cell(
        row=row_num,
        column=col_index["Best Retailer"],
    )

    best_retailer_cell.value = (
        f'=IFERROR('
        f'INDEX('
        f'{retailer_header_range},'
        f'MATCH('
        f'{best_price_col}{row_num},'
        f'{retailer_range},'
        f'0'
        f')'
        f'),'
        f'"N/A"'
        f')'
    )


# ============================================================
# CREATE / OPEN OUTPUT WORKBOOK
# ============================================================

def get_or_create_workbook(
    models: List[str],
    output_path: str = OUTPUT_FILE,
):
    """
    Load price_comparison.xlsx if it exists.

    Otherwise create it.

    Also adds new models that are not already present.
    """

    try:
        workbook = (
            openpyxl.load_workbook(
                output_path
            )
        )

        if "Price Comparison" in workbook.sheetnames:
            worksheet = workbook[
                "Price Comparison"
            ]

            # Keep an existing output workbook aligned with ALL_COLUMNS.
            if worksheet.max_column > len(ALL_COLUMNS):
                worksheet.delete_cols(
                    len(ALL_COLUMNS) + 1,
                    worksheet.max_column - len(ALL_COLUMNS),
                )

            for column, heading in enumerate(
                ALL_COLUMNS,
                start=1,
            ):
                worksheet.cell(
                    row=1,
                    column=column,
                    value=heading,
                )

            _style_header(worksheet)
        else:
            worksheet = (
                workbook.create_sheet(
                    "Price Comparison"
                )
            )

            worksheet.append(
                ALL_COLUMNS
            )

            _style_header(
                worksheet
            )

        existing_models = {}

        for row in range(
            2,
            worksheet.max_row + 1
        ):
            value = worksheet.cell(
                row=row,
                column=1,
            ).value

            if value is not None:
                existing_models[
                    str(value).strip()
                ] = row

        for model in models:

            if model in existing_models:
                continue

            new_row = (
                worksheet.max_row + 1
            )

            worksheet.cell(
                row=new_row,
                column=1,
                value=model,
            )

            _add_formulas(
                worksheet,
                new_row
            )

    except FileNotFoundError:

        workbook = (
            openpyxl.Workbook()
        )

        worksheet = workbook.active

        worksheet.title = (
            "Price Comparison"
        )

        worksheet.append(
            ALL_COLUMNS
        )

        _style_header(
            worksheet
        )

        for row_num, model in enumerate(
            models,
            start=2,
        ):
            worksheet.cell(
                row=row_num,
                column=1,
                value=model,
            )

            _add_formulas(
                worksheet,
                row_num,
            )

    # Set column widths
    for index, name in enumerate(
        ALL_COLUMNS,
        start=1,
    ):
        worksheet.column_dimensions[
            get_column_letter(index)
        ].width = max(
            20,
            len(name) + 4,
        )

    return (
        workbook,
        worksheet,
    )


# ============================================================
# WRITE RETAILER RESULT
# ============================================================

def update_retailer_cell(
    ws,
    retailer: str,
    model: str,
    value,
    url: Optional[str] = None,
):
    """
    Write a retailer result into the matching model row.

    Numeric prices remain numeric so Excel MIN() works.

    A URL can also be attached as a hyperlink.
    """

    col_index = {
        name: index + 1
        for index, name
        in enumerate(ALL_COLUMNS)
    }

    if retailer not in col_index:
        raise ValueError(
            f"Unknown retailer: {retailer}. "
            f"Add it to RETAILER_COLUMNS."
        )

    retailer_col = (
        col_index[retailer]
    )

    target_row = None

    for row in range(
        2,
        ws.max_row + 1
    ):
        existing_model = ws.cell(
            row=row,
            column=1,
        ).value

        if existing_model is None:
            continue

        if (
            str(existing_model).strip()
            == str(model).strip()
        ):
            target_row = row
            break

    # Add model if it is missing
    if target_row is None:

        target_row = (
            ws.max_row + 1
        )

        ws.cell(
            row=target_row,
            column=1,
            value=model,
        )

        _add_formulas(
            ws,
            target_row
        )

    cell = ws.cell(
        row=target_row,
        column=retailer_col,
    )

    cell.value = value

    # Price MUST remain numeric
    if isinstance(
        value,
        (int, float)
    ):
        cell.number_format = (
            '£#,##0.00'
        )

    else:
        cell.number_format = (
            "General"
        )

    # Add clickable product/search URL
    if url:
        cell.hyperlink = url

        cell.font = Font(
            color="1155CC",
            underline="single",
        )


# ============================================================
# LEGACY REQUESTS RETAILER RUNNER
#
# Keep this for retailers where requests / BeautifulSoup work.
#
# We will NOT use this function for Harvey Norman once we move
# Harvey Norman to Playwright.
# ============================================================

def run_retailer(
    retailer_name: str,
    search_url_template: str,
    parse_fn,
    input_path: str = INPUT_FILE,
    output_path: str = OUTPUT_FILE,
):
    """
    Generic retailer runner using requests + BeautifulSoup.

    Harvey Norman should use its own Playwright runner instead.
    """

    models = read_models(
        input_path
    )

    print(
        f"Found {len(models)} models "
        f"in {input_path}\n"
    )

    workbook, worksheet = (
        get_or_create_workbook(
            models,
            output_path,
        )
    )

    session = requests.Session()

    session.trust_env = False

    for index, model in enumerate(
        models,
        start=1,
    ):

        # Import locally because some retailer files
        # may still use this runner.
        from urllib.parse import quote_plus

        search_url = (
            search_url_template.format(
                q=quote_plus(model)
            )
        )

        print(
            f"[{index}/{len(models)}] "
            f"{retailer_name}: {model}"
        )

        policy = robots_status(
            search_url,
            session
        )

        if policy == "DISALLOWED":

            print(
                "    -> disallowed by robots.txt"
            )

            update_retailer_cell(
                worksheet,
                retailer_name,
                model,
                "BLOCKED_ROBOTS_TXT",
                url=search_url,
            )

            workbook.save(
                output_path
            )

            continue

        if policy == "UNAVAILABLE":

            print(
                "    -> robots.txt check failed"
            )

            update_retailer_cell(
                worksheet,
                retailer_name,
                model,
                "ROBOTS_CHECK_FAILED",
                url=search_url,
            )

            workbook.save(
                output_path
            )

            continue

        soup, status = fetch(
            search_url,
            session
        )

        polite_sleep()

        if status == "BLOCKED_BOT_DETECTION":

            update_retailer_cell(
                worksheet,
                retailer_name,
                model,
                "BLOCKED_BOT_DETECTION",
                url=search_url,
            )

            workbook.save(
                output_path
            )

            continue

        if (
            status == "ERROR"
            or soup is None
        ):

            update_retailer_cell(
                worksheet,
                retailer_name,
                model,
                "ERROR",
                url=search_url,
            )

            workbook.save(
                output_path
            )

            continue

        candidates = parse_fn(
            soup,
            model
        )

        chosen = best_match(
            candidates
        )

        if chosen is None:

            update_retailer_cell(
                worksheet,
                retailer_name,
                model,
                "NOT_FOUND",
                url=search_url,
            )

        elif (
            chosen.availability
            == "OUT_OF_STOCK"
        ):

            update_retailer_cell(
                worksheet,
                retailer_name,
                model,
                "OUT_OF_STOCK",
                url=(
                    chosen.url
                    or search_url
                ),
            )

        elif chosen.price is None:

            update_retailer_cell(
                worksheet,
                retailer_name,
                model,
                "PRICE_NOT_FOUND",
                url=(
                    chosen.url
                    or search_url
                ),
            )

        else:

            update_retailer_cell(
                worksheet,
                retailer_name,
                model,
                chosen.price,
                url=chosen.url,
            )

        # Save after each model
        workbook.save(
            output_path
        )

    workbook.save(
        output_path
    )

    print(
        f"\nDone. "
        f"{retailer_name} column updated "
        f"in {output_path}"
    )
